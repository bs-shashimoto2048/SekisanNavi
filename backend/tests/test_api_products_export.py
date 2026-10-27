"""`GET /api/products/{product_no}/export` (Excel出力) のAPI結合テスト。

指示: 「図面情報パネルに現在表示する盤情報と、アプリが保持している他の積算
仕様情報を、製番単位でExcelへ出力する」「旧図面情報パネルに表示していた
情報はExcel出力に含めない」「実際に取得・保持している情報だけを出力し、
未取得の値や仕様は推測で埋めない」「Excel出力は読み取り専用とし、既存DBや
積算データを変更しない」。

既存の`test_api_estimate_results.py::_setup_product_and_rule`と同じ方式
(product_df.csv/estcode_df.csv を書き込み、実際にBBox→評価器実行で
EstimateResultを作る)で、推測データではなく実際にアプリが生成・保持する
行を使って出力内容を検証する。
"""
import io
import sqlite3

import openpyxl

from app import config
from app.db.connection import get_connection
from app.domain.estimate_rules import JudgmentMethod, JudgmentScope, QuantityMethod, StandardCondition
from app.repositories.estimate_rule_masters import create_rule_master

_PRODUCT_DF_HEADER = (
    "BAN_MENNO,BAN_NO,PAGE,ZUMEI,BAN_MEISYOU,BAN_TYPE,BAN_H1,BAN_H2,BAN_W,BAN_D,"
    "KITEN_X,KITEN_Y,DETECT_AREA_X,DETECT_AREA_Y,FRAME_ORG_X,FRAME_ORG_Y,"
    "FRAME_MINI_X,FRAME_MINI_Y,SCALE_X,SCALE_Y"
)
_PANEL_1_1_ROW = "1,1.0,16,外形図,盤A,正面図,2300,2000,900,2200,100,100,900,2300,15990,11430,100,100,10,10"

_ESTCODE_DF_HEADER = (
    "MODEL,BAN_MENNO,BAN_NO,BAN_MEISYOU,BAN_H,BAN_W,BAN_D,BAN_CONNECT,PANEL,TRANS,"
    "IN_PANEL,SHIELD,DOOR_FRONT,DOOR_BACK,DOOR_STACK,DOOR_SIDE,DOOR_SMALL,FAN_ROOF,"
    "FAN_DOOR,MAIN_LINE,WIRE_MESH,STACK_PLATE,DRAWER_DEVICE,VCT_STAND,BUS_DUCT,"
    "PASSAGE,INPUT_CU_COEFF,SORT_ORDER"
)
_ESTCODE_ROW_1_1 = (
    "IS2,1,1.0,高圧受電盤,2300,900,2200,箱･左右(R),0,0,0,0,0,0,0,0,0,0,0,0,0,0,0,0,0,0,0.0,5"
)


def _write_cp932_csv(path, header: str, rows: list[str]):
    content = "\n".join([header, *rows]) + "\n"
    path.write_bytes(content.encode("cp932"))


def _configure_root(client, monkeypatch, root):
    monkeypatch.setattr(config, "ADMIN_PASSWORD", "test-admin-pass")
    res = client.put(
        "/api/settings/data-source",
        json={"root": str(root), "admin_password": "test-admin-pass"},
    )
    assert res.status_code == 200


def _page16_id(client) -> int:
    pages = client.get("/api/drawing-pages").json()
    return next(p["id"] for p in pages if p["page_no"] == 16)


def _first_master_item(client) -> dict:
    return client.get("/api/master-items").json()[0]


def _create_manual_detection(client, **overrides) -> dict:
    page_id = _page16_id(client)
    master_item = _first_master_item(client)
    body = {
        "drawing_page_id": page_id,
        "master_item_id": master_item["id"],
        "bbox_x": 0.1,
        "bbox_y": 0.1,
        "bbox_w": 0.05,
        "bbox_h": 0.05,
        **overrides,
    }
    res = client.post("/api/detections", json=body)
    assert res.status_code == 201
    return res.json()


def _set_evidence_type(db_path, detection_id: int, key: str) -> None:
    conn = sqlite3.connect(db_path)
    conn.execute("UPDATE detections SET evidence_type_key = ? WHERE id = ?", (key, detection_id))
    conn.commit()
    conn.close()


def _setup_product_with_result(client, monkeypatch, tmp_path, db_path):
    """`A1GV2421`製番に、盤情報(estcode_df 1件)+積算結果(評価器実行で1件)を用意する。"""
    product = tmp_path / "A1GV2421"
    product.mkdir()
    _write_cp932_csv(product / "product_df.csv", _PRODUCT_DF_HEADER, [_PANEL_1_1_ROW])
    _write_cp932_csv(product / "estcode_df.csv", _ESTCODE_DF_HEADER, [_ESTCODE_ROW_1_1])
    _configure_root(client, monkeypatch, tmp_path)

    detection = _create_manual_detection(client)
    _set_evidence_type(db_path, detection["id"], "side_door")

    master_item = _first_master_item(client)
    with get_connection(db_path) as conn:
        create_rule_master(
            conn,
            master_item_id=master_item["id"],
            judgment_method=JudgmentMethod.DRAWING_JUDGMENT,
            judgment_scope=JudgmentScope.PANEL,
            quantity_method=QuantityMethod.PER_EVIDENCE,
            judgment_condition=StandardCondition(required_evidence_types=["side_door"]),
        )

    evaluated = client.post("/api/products/A1GV2421/estimate-results/evaluate")
    assert evaluated.status_code == 200
    assert len(evaluated.json()["results"]) == 1
    return master_item


def _load_workbook(res) -> openpyxl.Workbook:
    assert res.status_code == 200
    return openpyxl.load_workbook(io.BytesIO(res.content))


def test_export_returns_xlsx_with_expected_headers_and_content_type(client, monkeypatch, tmp_path, db_path):
    master_item = _setup_product_with_result(client, monkeypatch, tmp_path, db_path)

    res = client.get("/api/products/A1GV2421/export")
    assert res.status_code == 200
    assert res.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in res.headers["content-disposition"]
    assert "A1GV2421" in res.headers["content-disposition"]

    wb = _load_workbook(res)
    assert wb.sheetnames == ["概要", "盤情報", "積算仕様"]

    overview = wb["概要"]
    assert overview["A1"].value == "項目"
    assert overview["A2"].value == "製番"
    assert overview["B2"].value == "A1GV2421"

    panel_sheet = wb["盤情報"]
    assert [c.value for c in panel_sheet[1]] == ["面番号", "盤番号", "盤名称", "型式", "高さ", "幅", "奥行", "接続"]
    panel_row = [c.value for c in panel_sheet[2]]
    assert panel_row == [1, 1, "高圧受電盤", "IS2", 2300, 900, 2200, "箱･左右(R)"]

    spec_sheet = wb["積算仕様"]
    assert [c.value for c in spec_sheet[1]] == [
        "対象区分",
        "面番号",
        "盤番号",
        "コード",
        "品名",
        "型式",
        "定格",
        "数量",
        "適用単位",
        "係数",
        "単価",
        "金額",
        "判定方法",
        "ステータス",
    ]
    spec_row = {header.value: cell.value for header, cell in zip(spec_sheet[1], spec_sheet[2])}
    # 面1/盤1向けのjudgment_scope=PANELルールで評価されたため「各盤」、面番号/盤番号が入る。
    assert spec_row["対象区分"] == "各盤"
    assert spec_row["面番号"] == 1
    assert spec_row["盤番号"] == 1
    assert spec_row["コード"] == master_item["code"]
    assert spec_row["品名"] == master_item["category"]
    assert spec_row["判定方法"] == "drawing_judgment"
    assert spec_row["ステータス"] == "auto"


def test_export_marks_missing_values_as_unconfirmed_not_fabricated(client, monkeypatch, tmp_path, db_path):
    # estcode_df.csv自体が存在しない(=盤情報を一切取得できない)製番。
    product = tmp_path / "A1NOINFO"
    product.mkdir()
    _configure_root(client, monkeypatch, tmp_path)

    res = client.get("/api/products/A1NOINFO/export")
    wb = _load_workbook(res)

    panel_sheet = wb["盤情報"]
    # ヘッダ行のみで、推測データの行を1件も作らない(指示: 未取得の値や仕様を
    # 推測で埋めない)。
    assert panel_sheet.max_row == 1

    spec_sheet = wb["積算仕様"]
    assert spec_sheet.max_row == 1


def test_export_distinguishes_product_wide_spec_from_panel_specific(client, monkeypatch, tmp_path, db_path):
    # 製番仕様(どの盤にも属さない)ルールのみを用意し、面番号/盤番号が空欄に
    # なることを確認する(指示: 製番仕様と各盤の仕様を混同しない)。
    product = tmp_path / "A1GV2421"
    product.mkdir()
    _write_cp932_csv(product / "product_df.csv", _PRODUCT_DF_HEADER, [_PANEL_1_1_ROW])
    _configure_root(client, monkeypatch, tmp_path)

    detection = _create_manual_detection(client)
    _set_evidence_type(db_path, detection["id"], "side_door")
    master_item = _first_master_item(client)
    with get_connection(db_path) as conn:
        create_rule_master(
            conn,
            master_item_id=master_item["id"],
            judgment_method=JudgmentMethod.DRAWING_JUDGMENT,
            judgment_scope=JudgmentScope.PRODUCT,
            quantity_method=QuantityMethod.PER_EVIDENCE,
            judgment_condition=StandardCondition(required_evidence_types=["side_door"]),
        )
    evaluated = client.post("/api/products/A1GV2421/estimate-results/evaluate")
    assert len(evaluated.json()["results"]) == 1

    res = client.get("/api/products/A1GV2421/export")
    wb = _load_workbook(res)
    spec_sheet = wb["積算仕様"]
    spec_row = {header.value: cell.value for header, cell in zip(spec_sheet[1], spec_sheet[2])}
    assert spec_row["対象区分"] == "製番仕様"
    # 製番仕様行の面番号/盤番号は空欄(None)にする(「未確認」ではなく、
    # 「このルールはどの盤にも属さない」という事実をそのまま表す)。
    assert spec_row["面番号"] is None
    assert spec_row["盤番号"] is None


def test_export_unknown_product_returns_404(client, monkeypatch, tmp_path):
    _configure_root(client, monkeypatch, tmp_path)
    res = client.get("/api/products/A1MISSING/export")
    assert res.status_code == 404


def test_export_does_not_modify_the_database(client, monkeypatch, tmp_path, db_path):
    _setup_product_with_result(client, monkeypatch, tmp_path, db_path)

    before = client.get("/api/products/A1GV2421/estimate-results").json()
    res = client.get("/api/products/A1GV2421/export")
    assert res.status_code == 200
    after = client.get("/api/products/A1GV2421/estimate-results").json()

    # 読み取り専用: エクスポートの前後でestimate_resultsの内容は一切変化しない。
    assert before == after
